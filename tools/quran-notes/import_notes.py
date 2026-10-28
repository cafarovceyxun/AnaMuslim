"""Doldurulmuş Excel faylındakı qeydləri `quran_translations_data.note`-a yazmaq üçün SQL hazırlayır.

    python3 tools/quran-notes/import_notes.py [hasiye-qeydleri.xlsx]

Bazaya özü yazmır — `out/notes_import.sql` yaradır və onu Oracle-da işlədən əmri çap edir.
Hər ayənin qeydi Excel-dəki sətirlərindən yığılır («(1) - …» formasında, nömrə sırası ilə) və
yalnız serverdəki hazırkı qeyddən **fərqlənəndə** yazılır. Hər UPDATE köhnə qeydi şərt kimi daşıyır:
fayl hazırlanandan sonra qeyd başqa yerdən dəyişibsə, həmin sətir toxunulmaz qalır (sayda görünür).

Qeyd tətbiqin redaktə növbəsindən (quran_edits) keçmir — birbaşa əsas cədvələ gedir, çünki bu,
adminin öz qeydləridir. Ehtiyat Frankfurt bazası gecəlik köçürmə ilə özü yenilənir.
"""
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

from common import ROOT, compose_note, fetch_translations, parse_note

SRC = sys.argv[1] if len(sys.argv) > 1 else str(ROOT / "docs/quran/hasiye-qeydleri.xlsx")
OUT = Path(__file__).resolve().parent / "out" / "notes_import.sql"
TAG = "$qn7$"

ws = load_workbook(SRC, data_only=True)["Qeydlər"]
header = [c.value for c in ws[1]]
if header[:7] != ["id", "Səhifə", "Sürə", "Ayə", "Nömrə", "Söz", "Qeyd"]:
    raise SystemExit(f"Gözlənilməyən sütunlar: {header} — make_sheet.py ilə yaradılmış fayl lazımdır.")

sheet = defaultdict(dict)
for row in ws.iter_rows(min_row=2, values_only=True):
    tid, _, _, ayah, num, _, note = row[:7]
    if tid is None:
        continue
    text = str(note).strip() if note is not None else ""
    if TAG in text:
        raise SystemExit(f"{ayah}: qeyddə {TAG} ardıcıllığı var — dəyişdirin.")
    sheet[int(tid)][int(num or 0)] = text

current = {r["id"]: r for r in fetch_translations()}
updates, skipped_unknown = [], []
for tid, pieces in sheet.items():
    r = current.get(tid)
    if r is None:
        skipped_unknown.append(tid)
        continue
    old_note = r["note"] or ""
    old_pieces = parse_note(old_note)
    # Excel-də sətri olmayan nömrə (məs. sonradan bazaya yazılmış) itməsin.
    merged = {**{k: v for k, v in old_pieces.items() if k not in pieces}, **pieces}
    new_note = compose_note(merged)
    if new_note == compose_note(old_pieces):
        continue
    updates.append((tid, r["chapter_no"], r["verse_no"], old_note, new_note))

if skipped_unknown:
    raise SystemExit(f"Serverdə olmayan id-lər: {skipped_unknown[:10]} — fayl köhnədir?")

if not updates:
    print("Dəyişiklik yoxdur — Excel-dəki qeydlər serverdəki ilə eynidir.")
    sys.exit()

OUT.parent.mkdir(exist_ok=True)
lines = ["begin;", "create temp table _notes_result(id bigint, ayah text, ok boolean) on commit drop;"]
for tid, c, v, old, new in updates:
    lines.append(
        f"with u as (update public.quran_translations_data set note = {TAG}{new}{TAG}, updated_at = now() "
        f"where id = {tid} and chapter_no = {c} and verse_no = {v} and coalesce(note, '') = {TAG}{old}{TAG} returning id) "
        f"insert into _notes_result select {tid}, '{c}:{v}', exists(select 1 from u);"
    )
lines += [
    "select 'yazıldı=' || count(*) filter (where ok) || ', toxunulmadı=' || count(*) filter (where not ok) "
    "|| coalesce(' (' || string_agg(ayah, ', ') filter (where not ok) || ')', '') from _notes_result;",
    "commit;",
]
OUT.write_text("\n".join(lines) + "\n")

for tid, c, v, old, new in updates:
    print(f"— {c}:{v}\n  köhnə: {old.strip()[:120] or '(yox)'}\n  yeni : {new[:120]}")
print(f"\n{len(updates)} ayənin qeydi dəyişəcək. SQL: {OUT}")
print("\nTətbiq etmək üçün:\n"
      f"ssh -i ~/.ssh/oracle_cafarovceyxun ubuntu@130.61.171.105 "
      f"'sudo docker exec -i supabase-db psql -U postgres -d postgres -v ON_ERROR_STOP=1 -q -t -A' < {OUT}")

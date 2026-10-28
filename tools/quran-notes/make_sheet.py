"""Haşiyə qeydləri üçün doldurulacaq Excel faylı yaradır.

    python3 tools/quran-notes/make_sheet.py [çıxış.xlsx]

Hər haşiyə nömrəsi bir sətirdir; «Qeyd» sütunu (sarı) doldurulur, sonra import_notes.py oxuyur.
Mövcud qeydlər öz sətirlərində əvvəlcədən yazılıb — redaktə etmək olar, silmək həmin nömrənin
qeydini silir. Gizli A sütunu (translation id) importun açarıdır, ona toxunulmamalıdır.
"""
import sqlite3
import sys
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from common import QURAN_DB, ROOT, compose_note, fetch_translations, markers, parse_note, surah_names

OUT = sys.argv[1] if len(sys.argv) > 1 else str(ROOT / "docs/quran/hasiye-qeydleri.xlsx")

rows = fetch_translations()
names = surah_names()
con = sqlite3.connect(QURAN_DB)
page = {}
for pg, a, b in con.execute(
    "select page_number, start_ayah_id, end_ayah_id from mushaf_map "
    "where mushaf_id = 1 and start_ayah_id is not null order by page_number"
):
    for i in range(a, b + 1):
        page.setdefault(i, pg)

items = []
for r in rows:
    ms = markers(r["text"])
    notes = parse_note(r["note"])
    if not ms and not notes:
        continue
    c, v = r["chapter_no"], r["verse_no"]
    nums = [n for n, _ in ms]
    entries = list(ms) + [(n, "— mətndə nömrə yoxdur") for n in sorted(notes) if n and n not in nums]
    if 0 in notes:
        entries.append((0, "— nömrəsiz qeyd"))
    for num, word in entries:
        items.append(dict(id=r["id"], page=page[c * 1000 + v], c=c, v=v, num=num, word=word,
                          note=notes.get(num, ""), text=(r["text"] or "").strip()))
items.sort(key=lambda x: (x["page"], x["c"], x["v"], x["num"]))

FONT = "Arial"
wb = Workbook()
ws = wb.active
ws.title = "Qeydlər"
headers = ["id", "Səhifə", "Sürə", "Ayə", "Nömrə", "Söz", "Qeyd", "Ayənin mətni"]
widths = [8, 8, 16, 8, 8, 18, 60, 90]
ws.append(headers)
for col, w in zip("ABCDEFGH", widths):
    ws.column_dimensions[col].width = w
ws.column_dimensions["A"].hidden = True
head_fill = PatternFill("solid", fgColor="1F5F3F")
for cell in ws[1]:
    cell.font = Font(name=FONT, bold=True, color="FFFFFF")
    cell.fill = head_fill
    cell.alignment = Alignment(vertical="center", wrap_text=True)

input_fill = PatternFill("solid", fgColor="FFF2CC")
done_fill = PatternFill("solid", fgColor="E2EFDA")
for it in items:
    ws.append([it["id"], it["page"], names[it["c"]], f"{it['c']}:{it['v']}", it["num"] or "", it["word"],
               it["note"], it["text"]])
    row = ws.max_row
    for cell in ws[row]:
        cell.font = Font(name=FONT)
        cell.alignment = Alignment(vertical="top", wrap_text=cell.column_letter in "GH")
    ws[f"G{row}"].fill = done_fill if it["note"] else input_fill
    ws[f"G{row}"].font = Font(name=FONT, color="0000FF")
    ws[f"H{row}"].font = Font(name=FONT, color="595959", size=9)
ws.freeze_panes = "G2"
ws.auto_filter.ref = f"B1:H{ws.max_row}"

legend = wb.create_sheet("Necə doldurmalı", 0)
legend.column_dimensions["A"].width = 110
lines = [
    ("Haşiyə qeydləri — doldurma qaydası", True),
    (f"Mətn serverdən {date.today().isoformat()}-də götürülüb: {len({i['id'] for i in items})} ayə, {len(items)} nömrə.", False),
    ("", False),
    ("1. «Qeydlər» vərəqində yalnız G sütununu («Qeyd») doldurun. Sarı xana — qeyd hələ yoxdur, yaşıl — artıq var.", False),
    ("2. Qeydin mətnini nömrəsiz yazın: «(1) -» hissəsini import özü əlavə edir.", False),
    ("   Nümunə (Bəqərə 2:2, nömrə 1): Təqva sahibi, Allahdan qorxub çəkinən, haramlardan özünü qoruyan birisi", False),
    ("3. Yaşıl xananı dəyişmək olar. Silsəniz, həmin nömrənin qeydi bazadan da silinir.", False),
    ("4. Boş saxladığınız sarı xanalar heç nəyi dəyişmir — faylı hissə-hissə doldurub import etmək olar.", False),
    ("5. Sətirləri silməyin, sıralamayın və gizli A sütununa (id) toxunmayın — import ayəni onunla tapır.", False),
    ("6. «Ayənin mətni» yalnız oxumaq üçündür; ayə mətnindəki dəyişiklik import olunmur.", False),
    ("", False),
    ("Bitirəndə faylı saxlayın və Claude-a deyin: «qeydləri import et».", True),
]
for text, bold in lines:
    legend.append([text])
    legend.cell(row=legend.max_row, column=1).font = Font(name=FONT, bold=bold, size=13 if bold else 11)

wb.active = 1
wb.save(OUT)
print(f"{OUT}: {len(items)} sətir, {len({i['id'] for i in items})} ayə, "
      f"qeydi olan {sum(1 for i in items if i['note'])} nömrə")
